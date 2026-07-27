"""Generate Excel + HTML + Markdown reports from the E2E results.json."""
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config


def _load():
    if not config.RESULTS_JSON.exists():
        return {"base_url": config.BASE_URL, "generated_at": "", "results": []}
    return json.loads(config.RESULTS_JSON.read_text(encoding="utf-8"))


def _stats(results):
    total = len(results)
    passed = sum(1 for r in results if r["status"] == "PASS")
    failed = sum(1 for r in results if r["status"] == "FAIL")
    skipped = sum(1 for r in results if r["status"] == "SKIP")
    pct = round(passed / total * 100, 2) if total else 0.0
    return total, passed, failed, skipped, pct


# --------------------------------------------------------------------------- #
def generate_excel(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    results = data["results"]
    total, passed, failed, skipped, pct = _stats(results)

    wb = Workbook()
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    yellow = PatternFill("solid", fgColor="FFEB9C")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center")

    # Sheet 1 — Summary
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Value"])
    style_header(ws, 2)
    rows = [
        ("Deployment URL", data.get("base_url", "")),
        ("Generated At", data.get("generated_at", "")),
        ("Total Tests", total),
        ("Passed", passed),
        ("Failed", failed),
        ("Skipped", skipped),
        ("Pass Percentage", f"{pct}%"),
    ]
    for r in rows:
        ws.append(list(r))
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60

    # Sheet 2 — All Results
    ws2 = wb.create_sheet("Test Results")
    cols = ["ID", "Name", "Category", "Page", "Status", "Duration (s)", "Expected", "Failure Reason", "Screenshot"]
    ws2.append(cols)
    style_header(ws2, len(cols))
    for r in results:
        ws2.append([r["id"], r["name"], r["category"], r.get("page", ""), r["status"],
                    r["duration"], str(r.get("expected", "")), r.get("reason", ""), r.get("screenshot", "")])
        fill = green if r["status"] == "PASS" else (red if r["status"] == "FAIL" else yellow)
        ws2.cell(row=ws2.max_row, column=5).fill = fill
    for col, w in zip("ABCDEFGHI", [10, 42, 12, 12, 8, 12, 14, 50, 30]):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"

    # Sheet 3 — Failures
    ws3 = wb.create_sheet("Failures")
    ws3.append(["ID", "Name", "Category", "Failure Reason", "Screenshot"])
    style_header(ws3, 5)
    for r in [x for x in results if x["status"] == "FAIL"]:
        ws3.append([r["id"], r["name"], r["category"], r.get("reason", ""), r.get("screenshot", "")])
    for col, w in zip("ABCDE", [10, 42, 12, 60, 30]):
        ws3.column_dimensions[col].width = w

    # Sheet 4 — By Category
    ws4 = wb.create_sheet("By Category")
    ws4.append(["Category", "Total", "Passed", "Failed", "Pass %"])
    style_header(ws4, 5)
    by_cat = defaultdict(lambda: [0, 0, 0])
    for r in results:
        by_cat[r["category"]][0] += 1
        if r["status"] == "PASS":
            by_cat[r["category"]][1] += 1
        elif r["status"] == "FAIL":
            by_cat[r["category"]][2] += 1
    for cat, (t, p, f) in sorted(by_cat.items()):
        ws4.append([cat, t, p, f, f"{round(p / t * 100, 1) if t else 0}%"])
    for col, w in zip("ABCDE", [16, 10, 10, 10, 10]):
        ws4.column_dimensions[col].width = w

    out = config.EXCEL_DIR / "Automation_Test_Report.xlsx"
    wb.save(out)
    return out


# --------------------------------------------------------------------------- #
def generate_html(data):
    results = data["results"]
    total, passed, failed, skipped, pct = _stats(results)
    by_cat = Counter(r["category"] for r in results)

    def badge(status):
        color = {"PASS": "#1FA463", "FAIL": "#E23744", "SKIP": "#C9A227"}.get(status, "#888")
        return f'<span style="background:{color};color:#fff;padding:2px 8px;border-radius:10px;font-size:12px">{status}</span>'

    rows = []
    for r in results:
        shot = f'<a href="../{r["screenshot"]}">view</a>' if r.get("screenshot") else ""
        rows.append(
            f"<tr class='{r['status']}'><td>{r['id']}</td><td>{r['name']}</td>"
            f"<td>{r['category']}</td><td>{badge(r['status'])}</td><td>{r['duration']}</td>"
            f"<td>{r.get('reason','')}</td><td>{shot}</td></tr>"
        )

    cat_rows = "".join(f"<tr><td>{c}</td><td>{n}</td></tr>" for c, n in sorted(by_cat.items()))

    html = f"""<!doctype html><html><head><meta charset="utf-8">
<title>CyberGuard AI — Live E2E Report</title>
<style>
 body{{font-family:Inter,Arial,sans-serif;margin:24px;background:#0f1220;color:#e8e8f0}}
 h1{{margin-bottom:4px}} .sub{{color:#9aa0b5;margin-top:0}}
 .cards{{display:flex;gap:16px;margin:20px 0;flex-wrap:wrap}}
 .card{{background:#191d33;border:1px solid #2a2f4a;border-radius:12px;padding:16px 22px;min-width:120px}}
 .card b{{font-size:28px;display:block}}
 table{{border-collapse:collapse;width:100%;margin-top:12px;background:#161a2e}}
 th,td{{border:1px solid #2a2f4a;padding:6px 10px;font-size:13px;text-align:left}}
 th{{background:#1f2540}} tr.FAIL{{background:#2a1720}} tr.PASS{{background:#14231b}}
 a{{color:#4c8dff}}
</style></head><body>
<h1>CyberGuard AI — Live GitHub Pages E2E Report</h1>
<p class="sub">Target: <a href="{data.get('base_url','')}">{data.get('base_url','')}</a> &nbsp;|&nbsp; {data.get('generated_at','')}</p>
<div class="cards">
 <div class="card"><b>{total}</b>Total</div>
 <div class="card" style="color:#4ade80"><b>{passed}</b>Passed</div>
 <div class="card" style="color:#f87171"><b>{failed}</b>Failed</div>
 <div class="card" style="color:#fbbf24"><b>{skipped}</b>Skipped</div>
 <div class="card"><b>{pct}%</b>Pass rate</div>
</div>
<h3>By Category</h3>
<table><tr><th>Category</th><th>Count</th></tr>{cat_rows}</table>
<h3>All Tests</h3>
<table>
<tr><th>ID</th><th>Name</th><th>Category</th><th>Status</th><th>Duration</th><th>Reason</th><th>Shot</th></tr>
{''.join(rows)}
</table>
</body></html>"""
    out = config.HTML_DIR / "execution-report.html"
    out.write_text(html, encoding="utf-8")
    return out


# --------------------------------------------------------------------------- #
def generate_summary_md(data):
    results = data["results"]
    total, passed, failed, skipped, pct = _stats(results)
    failures = [r for r in results if r["status"] == "FAIL"]

    lines = [
        "# Live GitHub Pages E2E Test Summary",
        "",
        f"**Deployment URL:** {data.get('base_url','')}",
        "",
        f"- **Total Tests:** {total}",
        f"- **Passed:** {passed}",
        f"- **Failed:** {failed}",
        f"- **Skipped:** {skipped}",
        f"- **Pass Percentage:** {pct}%",
        "",
        "## Failed Tests",
        "",
    ]
    if failures:
        lines.append("| Test | Reason |")
        lines.append("|------|--------|")
        for r in failures[:100]:
            reason = (r.get("reason", "") or "").replace("|", "\\|")[:160]
            lines.append(f"| {r['id']} — {r['name']} | {reason} |")
    else:
        lines.append("None 🎉 — all executed tests passed.")
    lines.append("")

    out = config.SUMMARY_DIR / "summary.md"
    out.write_text("\n".join(lines), encoding="utf-8")
    return out


def generate_all():
    data = _load()
    xlsx = generate_excel(data)
    html = generate_html(data)
    md = generate_summary_md(data)
    total, passed, failed, skipped, pct = _stats(data["results"])
    print(f"Reports generated: {total} tests, {passed} passed, {failed} failed ({pct}%)")
    print(f"  - {xlsx}")
    print(f"  - {html}")
    print(f"  - {md}")
    return failed


if __name__ == "__main__":
    sys.exit(1 if generate_all() > 0 else 0)
